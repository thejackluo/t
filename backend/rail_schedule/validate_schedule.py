#!/usr/bin/env python3
"""Cross-check exported schedules against raw GTFS and sample the public V3 API."""
import argparse
import csv
import gzip
import io
import json
import time
import urllib.parse
import urllib.request
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import openpyxl

if __package__:
    from .paths import OUTPUT_DIR
    from .schedule_lookup import lookup
else:
    from paths import OUTPUT_DIR
    from schedule_lookup import lookup


def validate(gtfs, output, api=True):
    z = zipfile.ZipFile(gtfs)
    def raw(name):
        with z.open(name + '.txt') as f:
            yield from csv.DictReader(io.TextIOWrapper(f, encoding='utf-8-sig'))
    def exported(name):
        with gzip.open(output / (name + '.csv.gz'), 'rt', newline='', encoding='utf-8') as f:
            yield from csv.DictReader(f)
    def seconds(s):
        h, m, s = map(int, s.split(':'))
        return h*3600+m*60+s
    def day(s):
        return datetime.strptime(s, '%Y%m%d').date()
    info = next(raw('feed_info'))
    start, end = day(info['feed_start_date']), day(info['feed_end_date'])
    trips = {r['trip_id']: r for r in raw('trips')}
    stops = {r['stop_id']: r for r in raw('stops')}
    active_dates = {}
    weekdays = ('monday', 'tuesday', 'wednesday', 'thursday', 'friday', 'saturday', 'sunday')
    for r in raw('calendar'):
        dates = set()
        d = day(r['start_date'])
        while d <= day(r['end_date']):
            if r[weekdays[d.weekday()]] == '1':
                dates.add(d.isoformat())
            d += timedelta(days=1)
        active_dates[r['service_id']] = dates
    for r in raw('calendar_dates'):
        dates = active_dates.setdefault(r['service_id'], set())
        if r['exception_type'] == '1':
            dates.add(day(r['date']).isoformat())
        else:
            dates.discard(day(r['date']).isoformat())
    profiles = {r['Profile ID']: r for r in exported('profiles')}
    service_routes = defaultdict(set)
    for r in trips.values():
        service_routes[r['service_id']].add(r['route_id'])
    calendar = list(exported('service_calendar'))
    assert len(calendar) == ((end-start).days+1)*8
    for r in calendar:
        expected = {s for s, dates in active_dates.items() if r['Service date'] in dates and r['Route'] in service_routes[s]}
        actual = set(filter(None, r['Active service IDs'].split('; ')))
        assert actual == expected, ('Calendar mismatch', r, expected)
        assert profiles[r['Profile ID']]['Active service IDs'] == r['Active service IDs']
    raw_by_service = defaultdict(Counter)
    for r in raw('stop_times'):
        t = trips[r['trip_id']]
        if t['route_id'] not in {p['Route'] for p in profiles.values()}:
            continue
        signature = (r['trip_id'], r['stop_sequence'], r['stop_id'], r['arrival_time'], r['departure_time'], r['pickup_type'] or '0', r['drop_off_type'] or '0')
        raw_by_service[t['service_id']][signature] += 1
    actual_by_profile = defaultdict(Counter)
    groups = defaultdict(list)
    destinations = defaultdict(list)
    previous, previous_destination = {}, {}
    events = list(exported('departures'))
    for r in events:
        t = trips[r['Trip ID']]
        assert r['Route'] == t['route_id']
        assert not r['Route'].startswith('Shuttle-')
        signature = (r['Trip ID'], r['Stop sequence'], r['Platform stop ID'], r['Arrival service time'], r['Departure service time'], '0' if r['Can board'] == '1' else '1', '0' if r['Can alight'] == '1' else '1')
        actual_by_profile[r['Profile ID']][signature] += 1
        assert r['Station ID'] == (stops[r['Platform stop ID']]['parent_station'] or r['Platform stop ID'])
        assert int(r['Departure service seconds']) == seconds(r['Departure service time'])
        key = (r['Profile ID'], r['Station ID'], r['Direction ID'])
        groups[key].append(r)
        kd = key + (r['Destination'],)
        destinations[kd].append(r)
        if r['Can board'] == '1':
            value = int(r['Departure service seconds'])
            if key in previous:
                assert int(r['Gap route seconds']) == value - previous[key] >= 0
            else:
                assert r['Gap route seconds'] == ''
            if kd in previous_destination:
                assert int(r['Gap destination seconds']) == value - previous_destination[kd] >= 0
            else:
                assert r['Gap destination seconds'] == ''
            previous[key] = previous_destination[kd] = value
        else:
            assert not r['Gap route seconds'] and not r['Gap destination seconds']
    for pid, p in profiles.items():
        expected = Counter()
        for service in filter(None, p['Active service IDs'].split('; ')):
            expected.update({key: n for key, n in raw_by_service[service].items() if trips[key[0]]['route_id'] == p['Route']})
        assert actual_by_profile[pid] == expected, ('Missing/extra stop events', pid)
    no_service, arrival_only, daily_checked = 0, 0, 0
    for r in exported('station_hours'):
        es = groups.get((r['Profile ID'], r['Station ID'], r['Direction ID']), [])
        board = sorted(e['Departure service time'] for e in es if e['Can board'] == '1')
        assert int(r['Boardable departures']) == len(board)
        if es:
            assert seconds(r['First arrival service time']) == min(seconds(e['Arrival service time']) for e in es)
            assert seconds(r['Last arrival service time']) == max(seconds(e['Arrival service time']) for e in es)
        else:
            assert r['Status'] == 'No scheduled rail'
            no_service += 1
        if board:
            assert seconds(r['First boarding service time']) == min(map(seconds, board))
            assert seconds(r['Last boarding service time']) == max(map(seconds, board))
            gaps = sorted(int(e['Gap route seconds']) for e in es if e['Gap route seconds'])
            if gaps:
                assert abs(float(r['Min gap minutes'])-min(gaps)/60) < .00051
                assert abs(float(r['Max gap minutes'])-max(gaps)/60) < .00051
        else:
            assert not r['First boarding service time'] and not r['Last boarding service time']
            arrival_only += bool(es)
        daily_checked += 1
    for r in exported('destination_hours'):
        es = destinations[(r['Profile ID'], r['Station ID'], r['Direction ID'], r['Destination'])]
        board = [int(e['Departure service seconds']) for e in es if e['Can board'] == '1']
        assert len(board) == int(r['Boardable departures'])
        if board:
            assert seconds(r['First boarding service time']) == min(board)
            assert seconds(r['Last boarding service time']) == max(board)
    wb = openpyxl.load_workbook(output / 'MBTA_Rail_Scheduling.xlsx', read_only=True, data_only=True)
    manifest = json.loads((output / 'manifest.json').read_text())
    for name, count in manifest['sheet_rows'].items():
        assert wb[name].max_row == count + 1, (name, wb[name].max_row, count)
    calc = wb['Wait lookup']
    expected = lookup(output, calc['B2'].value, calc['B3'].value, calc['B4'].value, calc['B5'].value,
                      seconds(calc['B7'].value), calc['B6'].value)
    assert calc['B12'].value == expected['next_departure_service_time']
    assert calc['B13'].value == expected['scheduled_wait_minutes']
    assert lookup(output, '2026-12-25', 'Red', 'place-pktrm', '0', 28800)['status'] == 'Date/route outside published coverage'
    assert 'No further' in lookup(output, start.isoformat(), 'Red', 'place-pktrm', '0', 35*3600)['status']
    wb.close()
    api_samples = []
    if api:
        # Public calls: validate a regular day, holiday, and late-night pattern independently.
        for d, route, station, direction in [('2026-10-22', 'Red', 'place-pktrm', '0'), ('2026-11-26', 'Blue', 'place-gover', '1'), ('2026-10-23', 'Green-E', 'place-north', '1')]:
            if not start <= date.fromisoformat(d) <= end:
                continue
            params = {'filter[date]': d, 'filter[route]': route, 'filter[route_type]': '0,1', 'filter[stop]': station,
                      'filter[direction_id]': direction, 'page[limit]': 1000, 'sort': 'departure_time'}
            url = 'https://api-v3.mbta.com/schedules?' + urllib.parse.urlencode(params)
            schedules = []
            while url:
                with urllib.request.urlopen(url, timeout=90) as response:
                    payload = json.load(response)
                schedules.extend(payload['data'])
                url = payload.get('links', {}).get('next')
                if url:
                    url = urllib.parse.urljoin('https://api-v3.mbta.com', url)
                time.sleep(3.2)
            pid = next(r['Profile ID'] for r in calendar if r['Service date'] == d and r['Route'] == route)
            es = [r for r in groups[(pid, station, direction)] if r['Can board'] == '1']
            expected_ids = Counter((r['Trip ID'], r['Platform stop ID'], r['Stop sequence']) for r in es)
            board = [s for s in schedules if s['attributes']['pickup_type'] == 0]
            actual_ids = Counter((s['relationships']['trip']['data']['id'], s['relationships']['stop']['data']['id'], str(s['attributes']['stop_sequence'])) for s in board)
            assert actual_ids == expected_ids, ('API identity/count mismatch', d, route, len(board), len(es))
            api_times = {s['relationships']['trip']['data']['id']: datetime.fromisoformat(s['attributes']['departure_time']) for s in board}
            for e in es:
                scheduled = date.fromisoformat(d) + timedelta(days=seconds(e['Departure service time']) // 86400)
                actual = api_times[e['Trip ID']]
                assert actual.date() == scheduled
                assert actual.hour*3600+actual.minute*60+actual.second == seconds(e['Departure service time']) % 86400
            api_samples.append({'service_date': d, 'route': route, 'station': station, 'direction': direction,
                                'boardable_calls_matched': len(board), 'trip_platform_and_timestamps': 'match'})
    report = {'validated_utc': datetime.now(timezone.utc).isoformat(), 'source_sha256': manifest['sha256'],
              'calendar_rows_checked': len(calendar), 'departure_rows_checked_against_raw_gtfs': len(events),
              'daily_station_rows_checked': daily_checked, 'no_scheduled_rail_rows': no_service,
              'arrival_only_rows': arrival_only, 'workbook_sheet_row_counts': 'match',
              'wait_lookup_cached_result': expected, 'api_samples': api_samples, 'result': 'PASS'}
    (output / 'validation.json').write_text(json.dumps(report, indent=2) + '\n')
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--gtfs', required=True, type=Path)
    p.add_argument('--out', type=Path, default=OUTPUT_DIR)
    p.add_argument('--skip-api', action='store_true')
    args = p.parse_args()
    validate(args.gtfs, args.out, not args.skip_api)
